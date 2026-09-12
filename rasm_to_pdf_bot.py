import os
import io
import json
import time
import threading
import logging

import requests
import telebot
from telebot import types
from flask import Flask
from PIL import Image
import img2pdf

from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# ------------------- SOZLAMALAR -------------------

BOT_TOKEN = os.environ.get("BOT_TOKEN")
GROQ_API_KEY = os.environ.get("GROQ_API_KEY")

if not BOT_TOKEN:
    raise RuntimeError(
        "BOT_TOKEN topilmadi! Render'da Environment Variables bo'limiga "
        "BOT_TOKEN qo'shishni unutmang."
    )

if not GROQ_API_KEY:
    logging.warning(
        "DIQQAT: GROQ_API_KEY topilmadi! /prezentatsiya funksiyasi ishlamaydi."
    )

# --- Taqdimot funksiyasi sozlamalari ---
ADMIN_ID = int(os.environ.get("ADMIN_ID", "8583388989"))
PAYMENT_CARD = os.environ.get("PAYMENT_CARD", "9860 1701 0940 9913")
PRESENTATION_PRICE = os.environ.get("PRESENTATION_PRICE", "10 000 so'm")
MAX_SLIDES = 20

GROQ_API_URL = "https://api.groq.com/openai/v1/chat/completions"
GROQ_MODEL = "openai/gpt-oss-120b"

UNSPLASH_ACCESS_KEY = os.environ.get("UNSPLASH_ACCESS_KEY")

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

bot = telebot.TeleBot(BOT_TOKEN, parse_mode="HTML")

# Har bir foydalanuvchi uchun vaqtinchalik rasmlar ro'yxati (RAM ichida saqlanadi)
# Struktura: { user_id: [ image_bytes, image_bytes, ... ] }
user_images = {}

# Har bir foydalanuvchi uchun vaqtinchalik matn qatorlari ro'yxati
# Struktura: { user_id: [ "matn qatori 1", "matn qatori 2", ... ] }
user_texts = {}

# Taqdimot buyurtmasi uchun holat mashinasi
# Struktura: { user_id: {"state": "...", "topic": "...", "slides": int} }
presentation_orders = {}


# ------------------- YORDAMCHI FUNKSIYALAR -------------------

def get_user_images(user_id):
    return user_images.setdefault(user_id, [])


def clear_user_images(user_id):
    user_images[user_id] = []


def get_user_texts(user_id):
    return user_texts.setdefault(user_id, [])


def clear_user_texts(user_id):
    user_texts[user_id] = []


def convert_to_jpeg_bytes(file_bytes):
    """Har qanday formatdagi rasmni img2pdf tushunadigan JPEG formatga o'giradi."""
    image = Image.open(io.BytesIO(file_bytes))
    if image.mode in ("RGBA", "P"):
        image = image.convert("RGB")
    output = io.BytesIO()
    image.save(output, format="JPEG")
    return output.getvalue()


def get_presentation_order(user_id):
    return presentation_orders.setdefault(user_id, {"state": "idle"})


def reset_presentation_order(user_id):
    presentation_orders[user_id] = {"state": "idle"}


def fetch_wikipedia_context(topic):
    """Wikipedia'dan mavzu bo'yicha qisqacha ma'lumot (extract) oladi.
    Avval o'zbekcha, topilmasa yoki juda qisqa bo'lsa inglizcha Wikipedia'ga murojaat qiladi.
    """
    for lang in ("uz", "en", "ru"):
        try:
            search_resp = requests.get(
                f"https://{lang}.wikipedia.org/w/api.php",
                params={
                    "action": "query",
                    "list": "search",
                    "srsearch": topic,
                    "format": "json",
                    "srlimit": 1,
                },
                timeout=10,
                headers={"User-Agent": "PresentationBot/1.0"},
            )
            search_resp.raise_for_status()
            results = search_resp.json().get("query", {}).get("search", [])
            if not results:
                continue

            page_title = results[0]["title"]

            extract_resp = requests.get(
                f"https://{lang}.wikipedia.org/w/api.php",
                params={
                    "action": "query",
                    "prop": "extracts",
                    "explaintext": 1,
                    "exchars": 4000,
                    "titles": page_title,
                    "format": "json",
                },
                timeout=10,
                headers={"User-Agent": "PresentationBot/1.0"},
            )
            extract_resp.raise_for_status()
            pages = extract_resp.json().get("query", {}).get("pages", {})
            for page in pages.values():
                extract = page.get("extract", "").strip()
                if extract and len(extract) > 200:
                    logger.info("Wikipedia (%s) topildi: %s", lang, page_title)
                    return extract
        except Exception:
            logger.exception("Wikipedia'dan (%s) ma'lumot olishda xatolik", lang)
            continue

    return None


def generate_outline_with_ai(topic, slide_count):
    """Groq AI orqali taqdimot uchun slaydlar mazmunini JSON ko'rinishida oladi."""
    wiki_context = fetch_wikipedia_context(topic)

    system_prompt = (
        "Sen professional taqdimot (prezentatsiya) tuzuvchi ekspertsan. "
        "Foydalanuvchi berayotgan mavzu bo'yicha CHUQUR, MA'LUMOTGA BOY va "
        "professional taqdimot tarkibini tuzasan. Agar senga Wikipedia'dan "
        "ma'lumot berilsa, undagi haqiqiy faktlar, raqamlar va tafsilotlardan "
        "FOYDALAN — o'zingdan taxminiy ma'lumot to'qib chiqarma. Har bir "
        "slaydda kamida 4-6 ta to'liq, aniq va foydali fikr (bullet) bo'lishi "
        "shart, har bir fikr kamida 8-15 so'zdan iborat to'liq gap bo'lsin "
        "(faqat qisqa iboralar emas). Slaydlar bir-birini takrorlamasin, har "
        "biri mavzuning boshqa jihatini yoritsin (masalan: kirish, tarixi, "
        "sabablari, statistika/raqamlar, ta'siri, yechimlar, misollar, xulosa "
        "kabi turli qismlarga bo'l). Shuningdek, har bir slayd uchun shu slayd "
        "mazmuniga mos, ingliz tilida 1-3 so'zdan iborat rasm qidiruv so'zini "
        "(image_query) ham ber (masalan 'climate change flood').\n\n"
        "Javobni FAQAT quyidagi JSON formatda ber, boshqa hech qanday matn, "
        "izoh yoki markdown belgisi qo'shma:\n"
        '{"title": "Taqdimot sarlavhasi", "subtitle": "Qisqa kichik sarlavha", '
        '"slides": [{"title": "Slayd sarlavhasi", "bullets": '
        '["to\'liq fikr 1", "to\'liq fikr 2", "to\'liq fikr 3", "to\'liq fikr 4"], '
        '"image_query": "english keywords"}]}'
    )

    user_prompt_parts = [
        f"Mavzu: {topic}",
        f"ANIQ {slide_count} ta kontent slaydi bo'lishi SHART (title slayddan tashqari, "
        f"kamroq ham, ko'proq ham emas — aynan {slide_count} ta).",
        "Mazmun professional, chuqur va real ma'lumotlarga asoslangan bo'lsin.",
        "Javob o'zbek tilida bo'lsin (image_query'dan tashqari).",
    ]
    if wiki_context:
        user_prompt_parts.append(
            f"\nQuyida Wikipedia'dan olingan real ma'lumot — shundan foydalan:\n"
            f"---\n{wiki_context}\n---"
        )
    user_prompt = "\n".join(user_prompt_parts)

    payload = {
        "model": GROQ_MODEL,
        "messages": [
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt},
        ],
        "temperature": 0.6,
        "max_tokens": 6000,
        "response_format": {"type": "json_object"},
    }
    if "gpt-oss" in GROQ_MODEL:
        payload["reasoning_effort"] = "low"

    last_outline = None
    for attempt in range(2):
        response = requests.post(
            GROQ_API_URL,
            headers={
                "Authorization": f"Bearer {GROQ_API_KEY}",
                "Content-Type": "application/json",
            },
            json=payload,
            timeout=90,
        )
        if not response.ok:
            logger.error("Groq API xatosi (%s): %s", response.status_code, response.text)
        response.raise_for_status()
        data = response.json()
        content = data["choices"][0]["message"]["content"]
        outline = json.loads(content)
        last_outline = outline

        slides = outline.get("slides", [])
        if len(slides) >= slide_count:
            outline["slides"] = slides[:slide_count]
            return outline

        logger.warning(
            "AI %s ta slayd o'rniga %s ta qaytardi, qayta urinilmoqda...",
            slide_count, len(slides)
        )

    return last_outline


def fetch_unsplash_image(query):
    """Unsplash'dan mavzuga mos rasm qidirib, bytes ko'rinishida qaytaradi."""
    if not UNSPLASH_ACCESS_KEY or not query:
        return None
    try:
        search_resp = requests.get(
            "https://api.unsplash.com/search/photos",
            params={"query": query, "per_page": 1, "orientation": "landscape"},
            headers={"Authorization": f"Client-ID {UNSPLASH_ACCESS_KEY}"},
            timeout=15,
        )
        search_resp.raise_for_status()
        results = search_resp.json().get("results", [])
        if not results:
            return None

        image_url = results[0]["urls"]["regular"]
        image_resp = requests.get(image_url, timeout=15)
        image_resp.raise_for_status()
        return io.BytesIO(image_resp.content)
    except Exception:
        logger.exception("Unsplash'dan rasm olishda xatolik (davom etamiz, rasmsiz)")
        return None


def build_pptx(outline):
    """Berilgan outline (dict) asosida .pptx faylini yaratadi va bytes qaytaradi."""
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    ACCENT_COLOR = RGBColor(0x1F, 0x4E, 0x79)
    TEXT_COLOR = RGBColor(0x22, 0x22, 0x22)

    blank_layout = prs.slide_layouts[6]

    # --- Sarlavha slaydi ---
    title_slide = prs.slides.add_slide(blank_layout)
    bg = title_slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
    )
    bg.fill.solid()
    bg.fill.fore_color.rgb = ACCENT_COLOR
    bg.line.fill.background()
    bg.shadow.inherit = False

    title_box = title_slide.shapes.add_textbox(
        Inches(1), Inches(2.7), Inches(11.3), Inches(1.8)
    )
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = outline.get("title", "Taqdimot")
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    p.alignment = PP_ALIGN.CENTER

    subtitle = outline.get("subtitle")
    if subtitle:
        sub_box = title_slide.shapes.add_textbox(
            Inches(1.5), Inches(4.5), Inches(10.3), Inches(1)
        )
        sp = sub_box.text_frame.paragraphs[0]
        sp.text = subtitle
        sp.font.size = Pt(20)
        sp.font.color.rgb = RGBColor(0xE0, 0xE0, 0xE0)
        sp.alignment = PP_ALIGN.CENTER

    # --- Kontent slaydlari ---
    for slide_data in outline.get("slides", []):
        slide = prs.slides.add_slide(blank_layout)

        # Sarlavha
        title_box = slide.shapes.add_textbox(
            Inches(0.6), Inches(0.35), Inches(12.1), Inches(0.9)
        )
        tp = title_box.text_frame.paragraphs[0]
        tp.text = slide_data.get("title", "")
        tp.font.size = Pt(28)
        tp.font.bold = True
        tp.font.color.rgb = ACCENT_COLOR

        # Sarlavha ostidagi chiziq
        line = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(1.15), Inches(4), Pt(3)
        )
        line.fill.solid()
        line.fill.fore_color.rgb = ACCENT_COLOR
        line.line.fill.background()
        line.shadow.inherit = False

        # Rasm (agar topilsa, o'ng tomonda)
        image_stream = fetch_unsplash_image(slide_data.get("image_query"))
        text_width = Inches(11.9)
        if image_stream:
            try:
                slide.shapes.add_picture(
                    image_stream, Inches(8.7), Inches(1.5), width=Inches(4.0), height=Inches(5.5)
                )
                text_width = Inches(7.7)
            except Exception:
                logger.exception("Rasmni slaydga joylashtirishda xatolik")

        # Matn (bullet'lar)
        body_box = slide.shapes.add_textbox(
            Inches(0.6), Inches(1.5), text_width, Inches(5.6)
        )
        body_tf = body_box.text_frame
        body_tf.word_wrap = True

        bullets = slide_data.get("bullets", [])
        for i, bullet in enumerate(bullets):
            bp = body_tf.paragraphs[0] if i == 0 else body_tf.add_paragraph()
            bp.text = f"•  {bullet}"
            bp.font.size = Pt(17)
            bp.font.color.rgb = TEXT_COLOR
            bp.space_after = Pt(12)

    output = io.BytesIO()
    prs.save(output)
    output.seek(0)
    return output


# ------------------- BOT BUYRUQLARI -------------------

@bot.message_handler(commands=["start"])
def handle_start(message):
    clear_user_images(message.from_user.id)
    clear_user_texts(message.from_user.id)
    reset_presentation_order(message.from_user.id)
    text = (
        "👋 Salom! Men foydali konvertatsiya botiman.\n\n"
        "🖼 <b>Rasm → PDF</b>\n"
        "1. Menga bir nechta rasm yuboring.\n"
        "2. /pdf buyrug'ini bering — bitta PDF faylga birlashtirib beraman.\n\n"
        "📝 <b>Matn → Word / Excel</b>\n"
        "1. Menga oddiy matn xabar(lar) yuboring.\n"
        "2. /word — matnlarni Word (.docx) hujjatiga aylantiraman.\n"
        "3. /excel — matnlarni Excel (.xlsx) jadvaliga aylantiraman "
        "(agar qatorda vergul bo'lsa, ustunlarga bo'lib joylashtiraman).\n\n"
        "🎓 <b>AI Taqdimot (Prezentatsiya)</b>\n"
        f"/prezentatsiya — mavzuni yozing, men AI yordamida to'liq taqdimot "
        f"(.pptx) tayyorlab beraman. Narxi: {PRESENTATION_PRICE}.\n\n"
        "🗑 /clear — hammasini (rasm va matnlarni) tozalash."
    )
    bot.reply_to(message, text)


@bot.message_handler(commands=["clear"])
def handle_clear(message):
    clear_user_images(message.from_user.id)
    clear_user_texts(message.from_user.id)
    reset_presentation_order(message.from_user.id)
    bot.reply_to(message, "🗑 Tanlangan rasmlar va matnlar tozalandi.")


@bot.message_handler(commands=["bekor"])
def handle_bekor(message):
    reset_presentation_order(message.from_user.id)
    bot.reply_to(message, "❌ Joriy jarayon bekor qilindi.")


@bot.message_handler(commands=["prezentatsiya"])
def handle_prezentatsiya_start(message):
    if not GROQ_API_KEY:
        bot.reply_to(
            message,
            "⚠️ Hozircha bu funksiya sozlanmagan (AI kaliti yo'q). "
            "Keyinroq urinib ko'ring."
        )
        return

    user_id = message.from_user.id
    order = get_presentation_order(user_id)
    order["state"] = "awaiting_topic"

    bot.reply_to(
        message,
        "🎓 Ajoyib! Taqdimot mavzusini yozing.\n"
        "Masalan: <i>«Iqlim o'zgarishi va uning oqibatlari»</i>"
    )


@bot.message_handler(content_types=["photo"])
def handle_photo(message):
    user_id = message.from_user.id
    order = get_presentation_order(user_id)

    # Agar foydalanuvchi to'lov skrinshotini kutayotgan bo'lsak,
    # bu rasmni PDF uchun emas, to'lov isboti sifatida qabul qilamiz.
    if order.get("state") == "awaiting_payment_screenshot":
        handle_payment_screenshot(message, order)
        return

    file_id = message.photo[-1].file_id  # eng yuqori sifatdagi versiyasi
    file_info = bot.get_file(file_id)
    downloaded_file = bot.download_file(file_info.file_path)

    try:
        jpeg_bytes = convert_to_jpeg_bytes(downloaded_file)
    except Exception as e:
        logger.exception("Rasmni qayta ishlashda xatolik")
        bot.reply_to(message, "❌ Bu rasmni qayta ishlab bo'lmadi, boshqa rasm yuboring.")
        return

    images = get_user_images(user_id)
    images.append(jpeg_bytes)

    bot.reply_to(
        message,
        f"✅ Rasm qabul qilindi ({len(images)}-ta). Yana rasm yuboring yoki /pdf deb yozing."
    )


def handle_payment_screenshot(message, order):
    user_id = message.from_user.id
    order["state"] = "awaiting_approval"

    bot.reply_to(
        message,
        "✅ To'lov skrinshoti qabul qilindi! Admin tasdiqlashini kuting "
        "(odatda tez orada). Tasdiqlangach, taqdimotingiz avtomatik yuboriladi."
    )

    username = message.from_user.username
    user_label = f"@{username}" if username else f"ID: {user_id}"

    caption = (
        f"💳 <b>Yangi to'lov so'rovi</b>\n\n"
        f"👤 Foydalanuvchi: {user_label} (id: {user_id})\n"
        f"📚 Mavzu: {order.get('topic')}\n"
        f"📊 Slaydlar soni: {order.get('slides')}\n\n"
        f"To'lovni tekshirib, tasdiqlang yoki rad eting:"
    )

    markup = types.InlineKeyboardMarkup()
    markup.add(
        types.InlineKeyboardButton("✅ Tasdiqlash", callback_data=f"approve_{user_id}"),
        types.InlineKeyboardButton("❌ Rad etish", callback_data=f"reject_{user_id}"),
    )

    bot.send_photo(
        ADMIN_ID,
        message.photo[-1].file_id,
        caption=caption,
        reply_markup=markup,
    )


@bot.callback_query_handler(func=lambda call: call.data.startswith(("approve_", "reject_")))
def handle_admin_decision(call):
    if call.from_user.id != ADMIN_ID:
        bot.answer_callback_query(call.id, "Bu tugma faqat admin uchun.")
        return

    action, user_id_str = call.data.split("_", 1)
    user_id = int(user_id_str)
    order = get_presentation_order(user_id)

    if action == "approve":
        bot.answer_callback_query(call.id, "Tasdiqlandi, taqdimot tayyorlanmoqda...")
        bot.edit_message_caption(
            caption=call.message.caption + "\n\n✅ TASDIQLANDI",
            chat_id=call.message.chat.id,
            message_id=call.message.message_id,
        )
        generate_and_send_presentation(user_id, order)
    else:
        bot.answer_callback_query(call.id, "Rad etildi.")
        bot.edit_message_caption(
            caption=call.message.caption + "\n\n❌ RAD ETILDI",
            chat_id=call.message.chat.id,
            message_id=call.message.message_id,
        )
        bot.send_message(
            user_id,
            "❌ To'lovingiz tasdiqlanmadi. Iltimos, to'lovni tekshirib, "
            "qaytadan /prezentatsiya buyrug'ini yuboring."
        )
        reset_presentation_order(user_id)


def generate_and_send_presentation(user_id, order):
    wait_msg = bot.send_message(user_id, "⏳ AI taqdimotingizni tayyorlamoqda, biroz kuting...")

    try:
        outline = generate_outline_with_ai(order["topic"], order["slides"])
        pptx_file = build_pptx(outline)
        pptx_file.name = "taqdimot.pptx"
    except Exception:
        logger.exception("Taqdimot yaratishda xatolik")
        bot.edit_message_text(
            "❌ Taqdimot yaratishda xatolik yuz berdi. Admin bilan bog'laning.",
            chat_id=user_id,
            message_id=wait_msg.message_id,
        )
        reset_presentation_order(user_id)
        return

    bot.send_document(
        user_id,
        pptx_file,
        caption="✅ Taqdimotingiz tayyor! Omad tilaymiz 🎓"
    )
    bot.delete_message(user_id, wait_msg.message_id)
    reset_presentation_order(user_id)


@bot.message_handler(content_types=["document"])
def handle_document(message):
    # Agar foydalanuvchi rasmni "fayl" sifatida yuborsa (siqilmagan holda)
    mime = message.document.mime_type or ""
    if not mime.startswith("image/"):
        bot.reply_to(message, "⚠️ Faqat rasm fayllarini qabul qilaman.")
        return

    user_id = message.from_user.id
    file_info = bot.get_file(message.document.file_id)
    downloaded_file = bot.download_file(file_info.file_path)

    try:
        jpeg_bytes = convert_to_jpeg_bytes(downloaded_file)
    except Exception:
        logger.exception("Hujjat-rasmni qayta ishlashda xatolik")
        bot.reply_to(message, "❌ Bu faylni qayta ishlab bo'lmadi.")
        return

    images = get_user_images(user_id)
    images.append(jpeg_bytes)
    bot.reply_to(
        message,
        f"✅ Rasm qabul qilindi ({len(images)}-ta). Yana rasm yuboring yoki /pdf deb yozing."
    )


@bot.message_handler(commands=["pdf"])
def handle_pdf(message):
    user_id = message.from_user.id
    images = get_user_images(user_id)

    if not images:
        bot.reply_to(
            message,
            "⚠️ Hali birorta ham rasm yubormadingiz. Avval rasm(lar) yuboring."
        )
        return

    wait_msg = bot.reply_to(message, "⏳ PDF tayyorlanmoqda...")

    try:
        pdf_bytes = img2pdf.convert(images)
    except Exception:
        logger.exception("PDF yaratishda xatolik")
        bot.edit_message_text(
            "❌ PDF yaratishda xatolik yuz berdi. Qaytadan urinib ko'ring.",
            chat_id=message.chat.id,
            message_id=wait_msg.message_id,
        )
        return

    pdf_file = io.BytesIO(pdf_bytes)
    pdf_file.name = "natija.pdf"

    bot.send_document(message.chat.id, pdf_file, caption="✅ Mana sizning PDF faylingiz!")
    bot.delete_message(message.chat.id, wait_msg.message_id)

    clear_user_images(user_id)


@bot.message_handler(commands=["word"])
def handle_word(message):
    user_id = message.from_user.id
    texts = get_user_texts(user_id)

    if not texts:
        bot.reply_to(
            message,
            "⚠️ Hali birorta ham matn yubormadingiz. Avval matn(lar) yuboring."
        )
        return

    wait_msg = bot.reply_to(message, "⏳ Word hujjati tayyorlanmoqda...")

    try:
        doc = Document()
        for line in texts:
            doc.add_paragraph(line)

        output = io.BytesIO()
        doc.save(output)
        output.seek(0)
        output.name = "natija.docx"
    except Exception:
        logger.exception("Word yaratishda xatolik")
        bot.edit_message_text(
            "❌ Word hujjatini yaratishda xatolik yuz berdi.",
            chat_id=message.chat.id,
            message_id=wait_msg.message_id,
        )
        return

    bot.send_document(message.chat.id, output, caption="✅ Mana sizning Word hujjatingiz!")
    bot.delete_message(message.chat.id, wait_msg.message_id)

    clear_user_texts(user_id)


@bot.message_handler(commands=["excel"])
def handle_excel(message):
    user_id = message.from_user.id
    texts = get_user_texts(user_id)

    if not texts:
        bot.reply_to(
            message,
            "⚠️ Hali birorta ham matn yubormadingiz. Avval matn(lar) yuboring."
        )
        return

    wait_msg = bot.reply_to(message, "⏳ Excel jadvali tayyorlanmoqda...")

    try:
        wb = Workbook()
        ws = wb.active
        ws.title = "Malumotlar"

        for line in texts:
            # Agar qatorda vergul bo'lsa, ustunlarga bo'lib joylashtiramiz
            cells = [cell.strip() for cell in line.split(",")]
            ws.append(cells)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        output.name = "natija.xlsx"
    except Exception:
        logger.exception("Excel yaratishda xatolik")
        bot.edit_message_text(
            "❌ Excel jadvalini yaratishda xatolik yuz berdi.",
            chat_id=message.chat.id,
            message_id=wait_msg.message_id,
        )
        return

    bot.send_document(message.chat.id, output, caption="✅ Mana sizning Excel jadvalingiz!")
    bot.delete_message(message.chat.id, wait_msg.message_id)

    clear_user_texts(user_id)


@bot.message_handler(func=lambda m: True, content_types=["text"])
def handle_plain_text(message):
    user_id = message.from_user.id
    order = get_presentation_order(user_id)
    state = order.get("state")

    # --- Taqdimot oqimi: mavzuni kutmoqda ---
    if state == "awaiting_topic":
        order["topic"] = message.text.strip()
        order["state"] = "awaiting_slide_count"
        bot.reply_to(
            message,
            f"📊 Nechta slayd bo'lsin? (1 dan {MAX_SLIDES} tagacha raqam yozing)"
        )
        return

    # --- Taqdimot oqimi: slaydlar sonini kutmoqda ---
    if state == "awaiting_slide_count":
        text = message.text.strip()
        if not text.isdigit() or not (1 <= int(text) <= MAX_SLIDES):
            bot.reply_to(
                message,
                f"⚠️ Iltimos, 1 dan {MAX_SLIDES} tagacha bo'lgan raqam yuboring."
            )
            return

        order["slides"] = int(text)
        order["state"] = "awaiting_payment_screenshot"
        bot.reply_to(
            message,
            f"💳 Ajoyib! Narxi: <b>{PRESENTATION_PRICE}</b>.\n\n"
            f"Quyidagi karta raqamiga to'lovni amalga oshiring:\n"
            f"<code>{PAYMENT_CARD}</code>\n\n"
            "To'lovni amalga oshirgach, chekning (skrinshotning) rasmini shu yerga yuboring."
        )
        return

    # --- Taqdimot oqimi: to'lov skrinshotini kutmoqda, lekin matn yuborilgan ---
    if state == "awaiting_payment_screenshot":
        bot.reply_to(message, "📸 Iltimos, matn emas, to'lov chekining RASMINI yuboring.")
        return

    if state == "awaiting_approval":
        bot.reply_to(message, "⏳ To'lovingiz hali admin tomonidan tekshirilmoqda, biroz kuting.")
        return

    # --- Aks holda, oddiy matn: Word/Excel uchun saqlaymiz ---
    texts = get_user_texts(user_id)
    texts.append(message.text)

    bot.reply_to(
        message,
        f"📝 Matn qabul qilindi ({len(texts)}-qator).\n"
        "Yana matn yuboring, yoki /word (Word) / /excel (Excel) deb yozing."
    )


# ------------------- FLASK KEEP-ALIVE SERVER -------------------
# Render Web Service turi doim ochiq portni kutadi, shuning uchun
# Flask serverni alohida oqim (thread)da ishga tushiramiz,
# botning polling jarayoni esa asosiy oqimda ishlaydi.

app = Flask(__name__)


@app.route("/")
def index():
    return "Bot ishlayapti ✅"


def run_flask():
    port = int(os.environ.get("PORT", 10000))
    app.run(host="0.0.0.0", port=port)


# ------------------- ISHGA TUSHIRISH -------------------

def clear_telegram_queue():
    """Eski webhook/navbatni tozalab, 409 (Conflict) xatosining oldini oladi."""
    try:
        url = f"https://api.telegram.org/bot{BOT_TOKEN}/deleteWebhook"
        requests.get(url, params={"drop_pending_updates": "true"}, timeout=10)
        logger.info("Telegram navbati tozalandi.")
    except Exception:
        logger.exception("Telegram navbatini tozalashda xatolik (muhim emas, davom etamiz)")


if __name__ == "__main__":
    flask_thread = threading.Thread(target=run_flask, daemon=True)
    flask_thread.start()

    clear_telegram_queue()
    time.sleep(2)

    logger.info("Bot polling boshlandi...")

    while True:
        try:
            bot.infinity_polling(skip_pending=True, timeout=30, long_polling_timeout=30)
        except Exception:
            logger.exception(
                "Polling to'xtadi (ehtimol vaqtinchalik Conflict xatosi). "
                "5 soniyadan keyin qayta urinamiz..."
            )
            time.sleep(5)
